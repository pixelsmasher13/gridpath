#!/usr/bin/env python3
"""
Grade a SpreadsheetBench run produced by run.mjs.

    python3 eval/spreadsheetbench/grade.py <run-dir-or-name> [--dataset <dir>]
        [--ids a,b] [--no-recalc] [--verbose]

Scoring is a line-for-line port of the official evaluation.py
(RUCKBReasoning/SpreadsheetBench): openpyxl data_only values at the task's
answer_position, numbers rounded to 2 dp, datetimes to Excel serial days,
numeric-looking strings coerced to float, "" == None, and a type mismatch is
a miss. A task passes when every listed range matches in every test case
(hard restriction). Verified 400 has one test case per task, so hard == soft.

Recalc: like the official open_spreadsheet.py, both the output and the golden
file are round-tripped through LibreOffice headless (profile seeded to
"always recalculate on load") so formula cells carry fresh cached values.
Golden recalcs are cached under <dataset>/_recalc/. --no-recalc grades raw
cached values (only sensible if the outputs were produced by Excel).

Deviations from the official script, all needed to run the Verified set as
shipped and each logged when they fire:
  * ranges without a sheet prefix use the task's `answer_sheet` when present
    (the official script always takes the golden's first sheet; on Verified
    the two disagree for exactly one task, 13-1);
  * commas inside quoted sheet names don't split the range list;
  * whole-column ranges (A:G) expand to the golden sheet's used rows;
  * a golden cell whose LibreOffice recalc is an Excel error while the
    shipped golden carries a real cached value is compared against the
    cached (Excel) value.

Outputs <run>/results.json (per-task verdicts + summary) and
<run>/eval_official_format.json (the official per-task record shape).
"""

import argparse
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from statistics import median

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))

# --- official comparator (verbatim port) --------------------------------------


def datetime_to_float(dt):
    excel_start_date = datetime.datetime(1899, 12, 30)
    delta = dt - excel_start_date
    return delta.days + delta.seconds / 86400.0


def transform_value(v):
    if isinstance(v, (int, float)):
        v = round(float(v), 2)
    elif isinstance(v, datetime.time):
        v = str(v)[:-3]
    elif isinstance(v, datetime.datetime):
        v = round(datetime_to_float(v), 0)
    elif isinstance(v, str):
        try:
            v = round(float(v), 2)
        except ValueError:
            pass
    return v


def compare_cell_value(v1, v2):
    v1 = transform_value(v1)
    v2 = transform_value(v2)
    if (v1 == "" and v2 is None) or (v1 is None and v2 == ""):
        return True
    if (v1 == "" and v2 == "") or (v1 is None and v2 is None):
        return True
    if type(v1) != type(v2):
        return False
    return v1 == v2


def col_num2name(n):
    name = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        name = chr(65 + remainder) + name
    return name


def col_name2num(name):
    num = 0
    for c in name:
        num = num * 26 + (ord(c) - ord("A") + 1)
    return num


def split_cell(ref):
    col, row = "", ""
    for ch in ref:
        if ch.isdigit():
            row += ch
        elif ch != "$":
            col += ch
    return col, row


def generate_cell_names(range_str, max_row):
    """Expand 'A1:B3' (or 'A:G', deviation) into cell names, columns outer."""
    if ":" not in range_str:
        return [range_str.replace("$", "")]
    start, end = range_str.split(":")
    sc, sr = split_cell(start)
    ec, er = split_cell(end)
    if not sr or not er:  # whole-column range — deviation, see module doc
        sr = sr or "1"
        er = er or str(max(1, max_row))
    columns = [col_num2name(i) for i in range(col_name2num(sc), col_name2num(ec) + 1)]
    return [f"{col}{row}" for col in columns for row in range(int(sr), int(er) + 1)]


RANGE_TAIL = re.compile(r"\$?[A-Za-z]{1,3}\$?\d*(:\$?[A-Za-z]{1,3}\$?\d*)?'?\s*$")


def split_ranges(answer_position):
    """Official: answer_position.split(','). Deviation: re-join fragments that
    were cut inside a quoted sheet name ("'b2b, sez, de'!A5:V10")."""
    out, acc = [], ""
    for frag in answer_position.split(","):
        acc = frag if not acc else acc + "," + frag
        if RANGE_TAIL.search(acc):
            out.append(acc)
            acc = ""
    if acc:
        out.append(acc)
    return out


def resolve_sheet_and_range(item, idx, data, wb_gt, notes):
    item = item.strip()
    if "!" in item:
        sheet_name, cell_range = item.rsplit("!", 1)
    else:
        sheet_name = None
        cell_range = item
        answer_sheet = data.get("answer_sheet")
        if answer_sheet:
            # A comma-separated answer_sheet lists one sheet per range.
            parts = [s.strip() for s in answer_sheet.split(",")]
            cand = parts[idx] if len(parts) > 1 and idx < len(parts) else parts[0]
            if cand in wb_gt.sheetnames:
                sheet_name = cand
                if cand != wb_gt.sheetnames[0]:
                    notes.append(f"answer_sheet '{cand}' != first sheet '{wb_gt.sheetnames[0]}' (deviation)")
        if sheet_name is None:
            sheet_name = wb_gt.sheetnames[0]
    sheet_name = sheet_name.strip().lstrip("'").rstrip("'")
    cell_range = cell_range.strip().lstrip("'").rstrip("'")
    return sheet_name, cell_range


EXCEL_ERRORS = {"#VALUE!", "#N/A", "#REF!", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def cell_level_compare(wb_gt, wb_proc, sheet_name, cell_range, wb_gt_raw, notes):
    if sheet_name not in wb_proc.sheetnames:
        return False, f"worksheet '{sheet_name}' not found in output"
    ws_gt = wb_gt[sheet_name]
    ws_proc = wb_proc[sheet_name]
    ws_raw = wb_gt_raw[sheet_name] if wb_gt_raw is not None and sheet_name in wb_gt_raw.sheetnames else None
    for cell_name in generate_cell_names(cell_range, max(ws_gt.max_row, ws_proc.max_row)):
        v_gt = ws_gt[cell_name].value
        v_proc = ws_proc[cell_name].value
        # Deviation: when LibreOffice can't evaluate the GOLDEN's own formula
        # (recalc → error) but Excel's cached value in the shipped golden is
        # a real value, the Excel value is the ground truth — the official
        # pipeline originally ran on Excel/win32com. Log it so the count of
        # such cells is visible.
        if ws_raw is not None and isinstance(v_gt, str) and v_gt in EXCEL_ERRORS:
            v_raw = ws_raw[cell_name].value
            if not (isinstance(v_raw, str) and v_raw in EXCEL_ERRORS):
                notes.append(f"{sheet_name}!{cell_name}: golden recalc={v_gt} but Excel cache={v_raw!r}; using cache (deviation)")
                v_gt = v_raw
        if not compare_cell_value(v_gt, v_proc):
            return False, f"{sheet_name}!{cell_name}: expected {v_gt!r}, got {v_proc!r}"
    return True, ""


def compare_workbooks(gt_file, proc_file, data, notes, gt_raw_file=None):
    if not os.path.exists(proc_file):
        return False, "output file missing"
    try:
        wb_gt = openpyxl.load_workbook(filename=gt_file, data_only=True)
        wb_proc = openpyxl.load_workbook(filename=proc_file, data_only=True)
    except Exception as e:  # noqa: BLE001
        return False, f"load error: {e}"
    wb_gt_raw = None
    if gt_raw_file and gt_raw_file != gt_file:
        try:
            wb_gt_raw = openpyxl.load_workbook(filename=gt_raw_file, data_only=True)
        except Exception:  # noqa: BLE001
            wb_gt_raw = None
    for idx, item in enumerate(split_ranges(data["answer_position"])):
        sheet_name, cell_range = resolve_sheet_and_range(item, idx, data, wb_gt, notes)
        if sheet_name not in wb_gt.sheetnames:
            return False, f"golden has no sheet '{sheet_name}' (answer_position={data['answer_position']!r})"
        ok, msg = cell_level_compare(wb_gt, wb_proc, sheet_name, cell_range, wb_gt_raw, notes)
        if not ok:
            return False, msg
    return True, ""


# --- LibreOffice recalc (same recipe as eval/grade.mjs) -----------------------

RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
 <item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
 <item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def find_soffice():
    for c in [
        os.environ.get("SOFFICE_PATH"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        "/opt/homebrew/bin/soffice",
    ]:
        if c and os.path.exists(c):
            return c
    return shutil.which("soffice")


def recalc(src, dest, soffice):
    """Convert xlsx→xlsx through a throwaway profile that always recalculates
    on load. Returns True on success (dest written)."""
    if os.path.exists(dest) and os.path.getmtime(dest) >= os.path.getmtime(src):
        return True
    tmp = tempfile.mkdtemp(prefix="sb-recalc-")
    try:
        profile = os.path.join(tmp, "profile")
        outdir = os.path.join(tmp, "out")
        os.makedirs(os.path.join(profile, "user"))
        os.makedirs(outdir)
        with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w") as fp:
            fp.write(RECALC_XCU)
        try:
            res = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--norestore",
                    f"-env:UserInstallation=file://{profile}",
                    "--convert-to",
                    "xlsx",
                    "--outdir",
                    outdir,
                    src,
                ],
                capture_output=True,
                text=True,
                timeout=180,
            )
        except subprocess.TimeoutExpired:
            return False
        converted = os.path.join(outdir, os.path.basename(src))
        if res.returncode != 0 or not os.path.exists(converted):
            return False
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        shutil.copyfile(converted, dest)
        return True
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def find_golden(dataset_dir, d, k=1):
    folder = os.path.join(dataset_dir, d.get("spreadsheet_path") or f"spreadsheet/{d['id']}")
    for name in (f"{k}_{d['id']}_golden.xlsx", f"{k}_{d['id']}_answer.xlsx", "golden.xlsx", "answer.xlsx"):
        p = os.path.join(folder, name)
        if os.path.exists(p):
            return p
    return None


# --- main ----------------------------------------------------------------------


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("run", help="run dir, or a name under eval/spreadsheetbench/runs/")
    ap.add_argument("--dataset", default=os.path.join(HERE, "spreadsheetbench_verified_400"))
    ap.add_argument("--ids", default=None)
    ap.add_argument("--no-recalc", action="store_true")
    ap.add_argument("--verbose", "-v", action="store_true")
    args = ap.parse_args()

    run_dir = args.run if os.path.isdir(args.run) else os.path.join(HERE, "runs", args.run)
    if not os.path.isdir(run_dir):
        sys.exit(f"run dir not found: {run_dir}")
    dataset_dir = os.path.abspath(args.dataset)
    with open(os.path.join(dataset_dir, "dataset.json")) as fp:
        dataset = {str(d["id"]): {**d, "id": str(d["id"])} for d in json.load(fp)}

    only = set(args.ids.split(",")) if args.ids else None
    task_ids = sorted(
        d for d in os.listdir(run_dir)
        if not d.startswith("_") and os.path.isdir(os.path.join(run_dir, d)) and d in dataset
        and (only is None or d in only)
    )
    if not task_ids:
        sys.exit("no graded tasks in run dir")

    soffice = None if args.no_recalc else find_soffice()
    if not args.no_recalc and not soffice:
        print("! LibreOffice (soffice) not found — set SOFFICE_PATH. Grading raw cached values.")

    results = []
    official = []
    for tid in task_ids:
        d = dataset[tid]
        task_dir = os.path.join(run_dir, tid)
        out = os.path.join(task_dir, "output.xlsx")
        meta_path = os.path.join(task_dir, "meta.json")
        meta = json.load(open(meta_path)) if os.path.exists(meta_path) else None
        notes = []
        verdict = {"id": tid, "instruction_type": d["instruction_type"], "answer_position": d["answer_position"]}

        golden = find_golden(dataset_dir, d)
        if golden is None:
            ok, msg = False, "golden file missing"
        elif meta is None:
            ok, msg = False, "not run (no meta.json)"
        elif not os.path.exists(out):
            ok, msg = False, "output.xlsx missing"
        else:
            gt = golden
            proc = out
            try:
                openpyxl.load_workbook(out, read_only=True).close()
            except Exception as e:  # noqa: BLE001
                notes.append(f"raw output not readable by openpyxl: {str(e).splitlines()[0][:120]}")
            if soffice:
                gt_rc = os.path.join(dataset_dir, "_recalc", f"{tid}_golden.xlsx")
                if recalc(golden, gt_rc, soffice):
                    gt = gt_rc
                else:
                    notes.append("golden recalc failed; using raw golden")
                out_rc = os.path.join(task_dir, "output.recalc.xlsx")
                if recalc(out, out_rc, soffice):
                    proc = out_rc
                else:
                    notes.append("output recalc failed; using raw output")
            try:
                ok, msg = compare_workbooks(gt, proc, d, notes, gt_raw_file=golden)
            except Exception as e:  # noqa: BLE001 — official script also swallows
                ok, msg = False, f"comparator error: {e}"

        verdict.update(
            passed=bool(ok),
            detail=msg,
            notes=notes,
            duration_ms=meta.get("duration_ms") if meta else None,
            input_tokens=meta.get("input_tokens") if meta else None,
            output_tokens=meta.get("output_tokens") if meta else None,
            batches=meta.get("batches") if meta else None,
            timed_out=bool(meta.get("timed_out")) if meta else None,
            run_error=meta.get("error") if meta else "not run",
            model=meta.get("model") if meta else None,
            effort=meta.get("effort") if meta else None,
        )
        results.append(verdict)
        official.append(
            {
                "id": tid,
                "instruction_type": d["instruction_type"],
                "test_case_results": [int(bool(ok))],
                "soft_restriction": 1.0 if ok else 0.0,
                "hard_restriction": 1 if ok else 0,
            }
        )
        mark = "PASS" if ok else "FAIL"
        line = f"  {mark} {tid:<8} {d['instruction_type'][:5]}"
        if meta:
            line += f" {round(meta.get('duration_ms', 0) / 1000):>4}s"
        if not ok or args.verbose:
            line += f"  {msg}"
        if notes and args.verbose:
            line += f"  [{'; '.join(notes)}]"
        print(line)

    n = len(results)
    passed = sum(r["passed"] for r in results)
    by_type = {}
    for r in results:
        t = by_type.setdefault(r["instruction_type"], [0, 0])
        t[0] += r["passed"]
        t[1] += 1
    ran = [r for r in results if r["duration_ms"] is not None]
    summary = {
        "run": os.path.basename(run_dir.rstrip("/")),
        "dataset": os.path.basename(dataset_dir),
        "tasks": n,
        "passed": passed,
        "pass_rate": round(100.0 * passed / n, 2) if n else 0.0,
        "by_type": {k: {"passed": v[0], "total": v[1], "pass_rate": round(100.0 * v[0] / v[1], 2)} for k, v in by_type.items()},
        "not_run": sum(1 for r in results if r["run_error"] == "not run"),
        "run_errors": sum(1 for r in results if r["run_error"] and r["run_error"] != "not run"),
        "timed_out": sum(1 for r in results if r["timed_out"]),
        "median_duration_s": round(median(r["duration_ms"] for r in ran) / 1000, 1) if ran else None,
        "total_input_tokens": sum(r["input_tokens"] or 0 for r in ran),
        "total_output_tokens": sum(r["output_tokens"] or 0 for r in ran),
        "models": sorted({f"{r['model']}/{r['effort']}" for r in ran if r["model"]}),
        "recalc": bool(soffice),
        "graded_at": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    with open(os.path.join(run_dir, "results.json"), "w") as fp:
        json.dump({"summary": summary, "tasks": results}, fp, indent=2)
    with open(os.path.join(run_dir, "eval_official_format.json"), "w") as fp:
        json.dump(official, fp, indent=4)

    print()
    print(f"{summary['run']} on {summary['dataset']}: {passed}/{n} = {summary['pass_rate']}%")
    for k, v in summary["by_type"].items():
        print(f"  {k}: {v['passed']}/{v['total']} = {v['pass_rate']}%")
    print(
        f"  not run {summary['not_run']}, run errors {summary['run_errors']}, timed out {summary['timed_out']}, "
        f"median {summary['median_duration_s']}s, tokens in/out {summary['total_input_tokens']}/{summary['total_output_tokens']}, "
        f"models {', '.join(summary['models']) or '-'}"
    )
    print(f"  results: {os.path.relpath(os.path.join(run_dir, 'results.json'))}")


if __name__ == "__main__":
    main()
